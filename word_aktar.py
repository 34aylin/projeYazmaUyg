"""python-docx ile başvuru formunu Word (.docx) dosyasına çevirir.
~ ile başlayan tahmini değerler belgede kırmızı renkte gösterilir."""

import io
import re
from datetime import date

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from form_alanlari import ALAN_ETIKETLERI, TUM_ALANLAR

_KIRMIZI = RGBColor(0xE5, 0x39, 0x35)
_GRI = RGBColor(60, 60, 60)
_BASLIK_GRI = RGBColor(235, 235, 235)

_TILDE_RE = re.compile(r"(~[^,;\n]+)")


def _renkli_paragraf(para, metin: str, font_size: int = 9) -> None:
    """Metni ~ parçaları kırmızı olacak şekilde run'lara böler."""
    if not metin:
        run = para.add_run("—")
        run.font.size = Pt(font_size)
        return
    parcalar = _TILDE_RE.split(metin)
    for parca in parcalar:
        if not parca:
            continue
        run = para.add_run(parca)
        run.font.size = Pt(font_size)
        if parca.startswith("~"):
            run.font.color.rgb = _KIRMIZI
            run.font.bold = True


def _bolum_baslik(doc: Document, baslik: str) -> None:
    from lxml import etree
    para = doc.add_paragraph()
    run = para.add_run(f"  {baslik}")
    run.font.size = Pt(10)
    run.font.bold = True
    # Arka plan rengi (shading) — XML ile eklenir
    pPr = para._p.get_or_add_pPr()
    shd = etree.SubElement(pPr, qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "EBEBEB")
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(4)


def _alan_yaz(doc: Document, etiket: str, metin: str) -> None:
    # Etiket
    etiket_para = doc.add_paragraph()
    run = etiket_para.add_run(etiket)
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = _GRI
    etiket_para.paragraph_format.space_before = Pt(0)
    etiket_para.paragraph_format.space_after = Pt(1)

    # Çok satırlı içerik — her satırı ayrı paragraf olarak yaz
    satirlar = metin.splitlines() if metin else [""]
    for i, satir in enumerate(satirlar):
        icerik_para = doc.add_paragraph()
        icerik_para.paragraph_format.space_before = Pt(0)
        icerik_para.paragraph_format.space_after = Pt(0 if i < len(satirlar) - 1 else 4)
        _renkli_paragraf(icerik_para, satir)


def word_olustur(sonuc: dict) -> bytes:
    doc = Document()

    # Kenar boşlukları
    for section in doc.sections:
        section.top_margin = Pt(50)
        section.bottom_margin = Pt(50)
        section.left_margin = Pt(70)
        section.right_margin = Pt(70)

    # Başlık
    baslik = doc.add_heading("Teknopark Ar-Ge Proje Başvuru Formu – Taslak", level=1)
    baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER
    baslik.runs[0].font.size = Pt(14)

    tarih_para = doc.add_paragraph(f"Oluşturulma: {date.today().strftime('%d.%m.%Y')}")
    tarih_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tarih_para.runs[0].font.size = Pt(8)
    tarih_para.runs[0].font.color.rgb = RGBColor(130, 130, 130)

    alanlar = sonuc.get("alanlar", {})
    deg = sonuc.get("arge_uygunluk_degerlendirmesi", {})

    # Ar-Ge uygunluk değerlendirmesi
    if deg:
        puan = deg.get("puan", "")
        renk = {"Yüksek": RGBColor(0, 140, 70), "Orta": RGBColor(180, 100, 0),
                "Düşük": RGBColor(190, 30, 30)}.get(puan, RGBColor(0, 0, 0))
        uyg_para = doc.add_paragraph()
        uyg_para.paragraph_format.space_before = Pt(6)
        r = uyg_para.add_run(f"Ar-Ge Uygunluk Değerlendirmesi: {puan}")
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = renk
        if deg.get("gerekce"):
            gerekce_para = doc.add_paragraph()
            _renkli_paragraf(gerekce_para, deg["gerekce"])

    # Form alanları
    _bolum_baslik(doc, "Form Alanları")
    for alan in TUM_ALANLAR:
        deger = alanlar.get(alan)
        if deger is None or deger == "" or deger == []:
            continue
        etiket = ALAN_ETIKETLERI.get(alan, alan)
        deger_str = ", ".join(str(d) for d in deger) if isinstance(deger, list) else str(deger)
        _alan_yaz(doc, etiket, deger_str)

    # Eksikler + uyarılar
    eksikler = (sonuc.get("eksik_bilgiler") or []) + (sonuc.get("uyarilar") or [])
    if eksikler:
        _bolum_baslik(doc, "Eksikler ve Uyarılar")
        for i, m in enumerate(eksikler, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _renkli_paragraf(p, f"{i}. {m}")

    sorular = sonuc.get("sorular") or []
    if sorular:
        _bolum_baslik(doc, "Önerilen Sorular")
        for i, s in enumerate(sorular, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _renkli_paragraf(p, f"{i}. {s}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
